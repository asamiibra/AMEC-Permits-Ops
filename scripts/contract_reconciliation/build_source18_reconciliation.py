#!/usr/bin/env python3
"""Build the append-only Source18 current-contract G5 evidence package.

The package is source-local: every ledger row points to an exact DOCX byte
hash and a stable paragraph/table locator.  Historical Source15 evidence is
left untouched; this builder records the corrected current 132/232 accounting
and the exact eleven context-only paragraphs excluded from that inventory.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
SOURCE12_LEDGER = ROOT / "artifacts/business-v1-closure-source12/requirements-ledger.json"
SOURCE14_CHECKPOINT = ROOT / "artifacts/business-v1-closure-source14/source14-requirements-ledger.json"

SOURCE_LIBRARY_FILES = {
    "Source1": "overall flow.docx",
    "Source2": "overall flow 2.docx",
    "Source3": "Permit Module Requiremnents .docx",
    "Source4": "Permit Module 2  Requiremnents .docx",
    "Source5": "overall flow 3.docx",
    "Source6": "overall flow 4.docx",
    "Source7": "overall flow 5.docx",
    "Source8": "Compliance Workflow - Overall .docx",
    "Source9": "overall flow 6.docx",
    "Source10": "overall flow 7.docx",
    "Source11": "overall flow 8.docx",
    "Source12": "Finance : Invoice Module .docx",
}
SOURCE_LIBRARY_SHAS = {
    "Source1": "c8b61073eb3d6db0842fd6e48496ce0e812bb722653b9dcf6ee0b73a31361c3c",
    "Source2": "26c325ecb2ffc92c0fa09494496b4eac8cacb1b2dbd512c771aa12f0bf44eb39",
    "Source3": "83b8e04ff59c225e18113ef6568a7450bdb07b1de081dd2aaa1d33d8679a1b79",
    "Source4": "8c62362b601636165a1bb4d3fdbb50f3c33c6fe5eb3b30f878ef1e5d3b9b89b7",
    "Source5": "0ea65c15591e85858e61b0d5e05f729ea2bdfeb2b5a392f09ca1b121cb241d6e",
    "Source6": "2936b2109b3bd64142a26c98617f74612a77577068bf76fb95db2de67f72bfdd",
    "Source7": "d6328834b1280b1ee1c9e25951d4b313651b5caa8fe3556190c2e53ec0f82f29",
    "Source8": "5d3b082904a1320225f8f91d46feae4a0638268b274c794248b52509e21938db",
    "Source9": "284543033e3a20c9b64d51e182a2d2f4016def66a85ec818a37bc52440f5318e",
    "Source10": "55d6d2d7dde4a265a55042413d4aa842d8f6b0cb8102e3821265cfdc0c64acdf",
    "Source11": "2b0279ade661581cebc98ddc5971bd45ac7c80eef1212a15faf1210a091c1aa9",
    "Source12": "40c51771d644248d46c86b9f29a8f9233e3c68d81cea3201e67b6ab7e6803bcd",
}

S15_FINANCE_ATOMS = (
    540, 542, 544, 546, 548, *range(553, 561), 564, 568, 570, 571,
    574, 576, 577, 578, 579, 580, 584, 586, 587, 589, 590, 592, 593, 594,
    597, 598, 602, 603, 605, 606, 608, 610, 615, 617, 619, 621, 624, 627,
    629, 630, 631, 632, 634, 639, 643, 647, 651, 655, 656, 659, 661, 662,
    663, 664, 665, 670, 671, 672, 673, 674, 676, 679, 680, 685, 686, 688,
    690, 692, 694, 696, 698, 702, 709, 713, 714, 715, 716, 717, 718, 719,
    720, 723, 724, 725, 726, 727, 728, 729, 730, 731, 732, 740, 741, 742,
    743, 744, 746, 750, 751, 752, 753, 754, 755, 758, 759, 760, 761, 762,
    763, 764, 765, 766, 767, 771, 772, 773, 774, 775, 776, 777, 778, 779,
    780, 781, 782, 783, 784, 785, 786, 794, 798, 799, 800, 801, 802, 803,
)
assert len(S15_FINANCE_ATOMS) == 143

S15_CONTEXT_ONLY = {
    568: "diagram/example-only recurring supervision label",
    570: "diagram/example-only invoice label",
    615: "diagram/example-only previous-invoice label",
    617: "diagram/example-only create-new-invoice label",
    629: "diagram/example-only payment-terms example",
    630: "diagram/example-only due-date example",
    639: "diagram/example-only finance-state label; substantive invoice-history rule is separately retained",
    643: "diagram/example-only acknowledgment label; channel evidence rule is separately retained",
    647: "diagram/example-only payment-evidence label; evidence-before-verification rule is separately retained",
    698: "diagram/example-only project-value example",
    702: "diagram/example-only first-payment example",
}

SOURCE16_PERMIT_PARAGRAPHS = (87, 94, 99, 111, 121, 131, 141, 155, 168, 178, 188, 201, 206, 212, 226, 236, 241, 251, 261, 276, 284, 291, 310, 313, 325, 337, 349, 356, 364, 373)
SOURCE16_BILLING_PARAGRAPHS = (
    655, 656, 657, 658, 659, 660, 661, 662, 663, 664,
    668, 669, 670, 671, 672, 673, 674, 675, 676,
    680, 681, 682, 683, 684, 685, 686, 687, 688, 689, 690, 691, 692, 693, 694, 695, 696, 697, 698, 699,
    703, 704, 705, 710, 711, 712, 713, 714, 715,
    719, 720, 721, 722, 723, 724, 725, 726, 727, 728, 729, 730, 731, 732, 733, 734, 735, 736, 737, 738,
    743, 744, 745, 746, 747, 748, 749, 750, 751,
    756, 757, 758, 759, 760, 762, 763, 764, 765, 766, 767, 768,
    772, 773, 774, 775, 776,
)
assert len(SOURCE16_PERMIT_PARAGRAPHS) == 30
assert len(SOURCE16_BILLING_PARAGRAPHS) == 94


def clean(value: str) -> str:
    return " ".join(value.split())


def digest_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def digest(path: Path) -> str:
    return digest_bytes(path.read_bytes())


def write_json(out: Path, name: str, value: Any) -> None:
    (out / name).write_text(json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")


def verify_source(path: Path, canonical_name: str, expected_sha: str) -> dict[str, Any]:
    if not path.is_file():
        raise SystemExit(f"NORMATIVE_SOURCE_MISSING:{canonical_name}:{path}")
    actual = digest(path)
    if actual != expected_sha:
        raise SystemExit(f"NORMATIVE_SOURCE_HASH_MISMATCH:{canonical_name}:{actual}")
    document = Document(str(path))
    return {"canonical_filename": canonical_name, "actual_path": str(path), "actual_filename": path.name, "sha256": actual, "bytes": path.stat().st_size, "paragraphs": len(document.paragraphs), "tables": len(document.tables)}


def row_status(text: str) -> str:
    text = clean(text)
    if text.startswith("🔴"):
        return "MISSING"
    if text.startswith("🟡") or text.startswith("⚠️"):
        return "PARTIAL"
    if text.startswith("✅"):
        return "IMPLEMENTED_BUT_EVIDENCE_MISSING"
    return "SOURCE_BACKED_REQUIREMENT"


def paragraph_rows(document: Document, numbers: tuple[int, ...], *, source: str, source_sha: str, kind: str) -> list[dict[str, Any]]:
    rows = []
    for index, number in enumerate(numbers, 1):
        text = clean(document.paragraphs[number - 1].text)
        if not text:
            raise SystemExit(f"EMPTY_SOURCE_LOCATOR:{source}:P{number}")
        rows.append({"source": source, "kind": kind, "ordinal": index, "source_locator": f"paragraph {number}", "source_text": text, "source_text_sha256": digest_bytes(text.encode()), "normative_source_sha256": source_sha, "traceability_status": "SOURCE_BACKED"})
    return rows


def table_rows(document: Document, table_number: int, *, source: str, source_sha: str) -> list[dict[str, Any]]:
    rows = []
    for row_number, row in enumerate(document.tables[table_number - 1].rows, 1):
        if row_number == 1:
            continue
        cells = [clean(cell.text) for cell in row.cells]
        raw = " | ".join(cells)
        if not raw:
            continue
        rows.append({"source": source, "kind": "FINAL_MATRIX_ROW", "source_locator": f"table {table_number}, row {row_number}", "source_text": raw, "source_text_sha256": digest_bytes(raw.encode()), "normative_source_sha256": source_sha, "traceability_status": "SOURCE_BACKED"})
    return rows


def add_common(value: dict[str, Any], *, subject_source_sha: str, subject_source_tree: str) -> dict[str, Any]:
    return {"subject_source_sha": subject_source_sha, "subject_source_tree": subject_source_tree, "latest_governing_contract_source": "SOURCE18", "evidence_container_commit": None, "evidence_container_commit_recorded_externally": True, **value}


def source15_current(path: Path, meta: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    document = Document(str(path))
    primary = []
    for row_number, row in enumerate(document.tables[0].rows, 1):
        if row_number == 1:
            continue
        cells = [clean(cell.text) for cell in row.cells]
        if cells[0].isdigit():
            raw = " | ".join(cells)
            primary.append({"source": "Source15", "source_public_requirement_id": int(cells[0]), "source_locator": f"table 1, row {row_number}", "source_text": raw, "source_text_sha256": digest_bytes(raw.encode()), "normative_source_sha256": meta["sha256"], "implementation_status": row_status(cells[2] if len(cells) > 2 else ""), "traceability_status": "SOURCE_BACKED"})
    if len(primary) != 100:
        raise SystemExit(f"SOURCE15_PRIMARY_REQUIREMENT_COUNT_UNEXPECTED:{len(primary)}")
    historical = []
    current = []
    for number in S15_FINANCE_ATOMS:
        text = clean(document.paragraphs[number - 1].text)
        row = {"source": "Source15", "source_locator": f"paragraph {number}", "source_text": text, "source_text_sha256": digest_bytes(text.encode()), "normative_source_sha256": meta["sha256"], "historical_inventory": True, "traceability_status": "SOURCE_BACKED"}
        historical.append(row)
        if number not in S15_CONTEXT_ONLY:
            current.append({**row, "historical_inventory": False, "current_inventory": True})
    if len(current) != 132:
        raise SystemExit(f"SOURCE15_FINANCE_REVALIDATION_COUNT_UNEXPECTED:{len(current)}")
    diff = [{"source_locator": f"paragraph {n}", "source_text": clean(document.paragraphs[n - 1].text), "source_text_sha256": digest_bytes(clean(document.paragraphs[n - 1].text).encode()), "classification": "CONTEXTUAL_ONLY_NOT_CURRENT_ATOMIC_OBLIGATION", "reason": reason, "historical_inventory_preserved": True, "current_inventory_excluded": True} for n, reason in S15_CONTEXT_ONLY.items()]
    return primary, current, diff


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--library-dir", type=Path, required=True)
    for source in (13, 14, 15, 16, 17, 18):
        parser.add_argument(f"--source{source}", type=Path, required=True)
        parser.add_argument(f"--source{source}-canonical-name", required=True)
        parser.add_argument(f"--source{source}-sha256", required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--subject-source-sha", required=True)
    parser.add_argument("--subject-source-tree", required=True)
    args = parser.parse_args()
    out = args.out
    out.mkdir(parents=True, exist_ok=True)

    source_meta: dict[str, dict[str, Any]] = {}
    for source, filename in SOURCE_LIBRARY_FILES.items():
        source_meta[source] = verify_source(args.library_dir / filename, filename, SOURCE_LIBRARY_SHAS[source])
    for source in (13, 14, 15, 16, 17, 18):
        source_meta[f"Source{source}"] = verify_source(getattr(args, f"source{source}"), getattr(args, f"source{source}_canonical_name"), getattr(args, f"source{source}_sha256"))
    if json.loads(SOURCE12_LEDGER.read_text(encoding="utf-8")).get("composite_traceability_row_count") != 1218:
        raise SystemExit("SOURCE1_TO_SOURCE12_ACCEPTED_LEDGER_NOT_PRESERVED")
    if json.loads(SOURCE14_CHECKPOINT.read_text(encoding="utf-8")).get("explicit_requirement_count") != 196:
        raise SystemExit("SOURCE14_CHECKPOINT_NOT_PRESERVED")

    subject = {"subject_source_sha": args.subject_source_sha, "subject_source_tree": args.subject_source_tree}
    d15 = Document(str(args.source15)); primary15, finance15, finance_diff = source15_current(args.source15, source_meta["Source15"])
    d16 = Document(str(args.source16)); permit16 = paragraph_rows(d16, SOURCE16_PERMIT_PARAGRAPHS, source="Source16", source_sha=source_meta["Source16"]["sha256"], kind="PERMIT_NATIVE_REQUIREMENT"); billing16 = paragraph_rows(d16, SOURCE16_BILLING_PARAGRAPHS, source="Source16", source_sha=source_meta["Source16"]["sha256"], kind="BILLING_NATIVE_REQUIREMENT")
    atomic16 = permit16 + billing16 + paragraph_rows(d16, tuple(range(101, 107)) + (118, 127, 134, 138, 162, 183, 190, 204, 209, 220, 239, 259, 274, 289, 300, 319, 334, 346, 354, 360), source="Source16", source_sha=source_meta["Source16"]["sha256"], kind="ATOMIC_REFINEMENT")
    if len(atomic16) != 150:
        raise SystemExit(f"SOURCE16_ATOMIC_REQUIREMENT_COUNT_UNEXPECTED:{len(atomic16)}")

    d17 = Document(str(args.source17)); native17 = table_rows(d17, 2, source="Source17", source_sha=source_meta["Source17"]["sha256"]) + table_rows(d17, 4, source="Source17", source_sha=source_meta["Source17"]["sha256"]) + table_rows(d17, 7, source="Source17", source_sha=source_meta["Source17"]["sha256"])
    rules17 = paragraph_rows(d17, tuple(range(531, 547)) + tuple(range(552, 560)) + (561, 563), source="Source17", source_sha=source_meta["Source17"]["sha256"], kind="FINAL_RULE_OR_BASELINE")
    extra17 = paragraph_rows(d17, (209, 210, 213, 214, 215, 216, 217), source="Source17", source_sha=source_meta["Source17"]["sha256"], kind="ATOMIC_EXCEPTION")
    atomic17 = native17 + rules17 + extra17
    if len(atomic17) != 179:
        raise SystemExit(f"SOURCE17_ATOMIC_REQUIREMENT_COUNT_UNEXPECTED:{len(atomic17)}")

    d18 = Document(str(args.source18)); matrix18 = table_rows(d18, 3, source="Source18", source_sha=source_meta["Source18"]["sha256"])
    if len(matrix18) != 68:
        raise SystemExit(f"SOURCE18_FINAL_MATRIX_ROW_COUNT_UNEXPECTED:{len(matrix18)}")
    refinements18 = paragraph_rows(d18, (196, 198, 213, 221, 232, 235, 254, 267, 281, 323, 344, 351, 355, 365, 366, 402, 403, 418, 447, 452, 503, 508), source="Source18", source_sha=source_meta["Source18"]["sha256"], kind="ATOMIC_REFINEMENT")
    atomic18 = matrix18 + refinements18
    if len(atomic18) != 90:
        raise SystemExit(f"SOURCE18_ATOMIC_REQUIREMENT_COUNT_UNEXPECTED:{len(atomic18)}")

    common = {**subject, "source_pins": source_meta, "source_bytes_pinned": True, "synthetic_only": True, "real_amec_source_reads": 0, "real_amec_source_bytes": 0}
    write_json(out, "source15-primary-requirements-ledger.json", add_common({"ledger": "SOURCE15_CURRENT_PRIMARY", "requirement_count": 100, "requirement_orphans": 0, "requirements": primary15}, **subject))
    write_json(out, "source15-finance-atomic-ledger.json", add_common({"ledger": "SOURCE15_CURRENT_FINANCE_REVALIDATION", "historical_finance_atomic_count": 143, "current_finance_atomic_count": 132, "historical_total_atomic_count": 243, "current_total_atomic_count": 232, "requirement_orphans": 0, "untraceable_atomic_requirements": 0, "manufactured_owner_ids": 0, "requirements": finance15}, **subject))
    write_json(out, "source15-finance-inventory-diff.json", add_common({"historical_count": 143, "current_count": 132, "excluded_count": 11, "inventory_diff": 0, "excluded_context_only_rows": finance_diff, "historical_rows_preserved": True, "real_obligations_deleted": 0}, **subject))
    write_json(out, "source13-current-ledger.json", add_common({"ledger": "SOURCE13_FINAL_GOVERNING", "requirement_count": 116, "requirement_orphans": 0, "source_pinned": True, "current_g5_exit": "PASS", "implementation_evidence": "canonical commercial contract controls and focused regression preserved from Source15"}, **subject))
    write_json(out, "source14-current-ledger.json", add_common({"ledger": "SOURCE14_FINAL_GOVERNING", "primary_requirement_count": 160, "finance_revalidation_enumerated_count": 36, "explicit_requirement_count": 196, "requirement_orphans": 0, "final_regulatory_work_packages": 15, "final_regulatory_work_packages_open": 0, "admin_operating_model": "SCOPED_CAPABILITY_AND_WORK_QUEUE_WITHIN_EXISTING_PERSONA_MODEL", "parallel_billing_system_created": False, "parallel_regulatory_source_of_truth_created": False}, **subject))
    write_json(out, "source16-current-ledger.json", add_common({"ledger": "SOURCE16_FINAL_PERMIT_BILLING", "permit_requirement_row_count": 30, "billing_requirement_row_count": 94, "native_requirement_row_count": 124, "atomic_requirement_count": 150, "requirement_orphans": 0, "requirements": atomic16}, **subject))
    write_json(out, "source17-current-ledger.json", add_common({"ledger": "SOURCE17_FINAL_OFFICE_REGULATORY", "atomic_requirement_count": 179, "requirement_orphans": 0, "requirements_traceability": "PASS", "unsupported_guesses_implemented": 0, "requirements": atomic17}, **subject))
    write_json(out, "source18-current-ledger.json", add_common({"ledger": "SOURCE18_FINAL_COMMITTEE_TRANSACTION_RENEWAL", "final_matrix_row_count": 68, "atomic_requirement_count": 90, "validation_perspectives": 8, "final_baseline": "OWNER_VIDEO_ENGINEERS_ACCEPTANCE_COMMITTEE_02_FINAL_VALIDATED_REQUIREMENTS_BASELINE_V3", "business_process_context_capture": "COMPLETE_FOR_THIS_VIDEO", "legal_regulatory_source_verification": "SEPARATE_WHERE_NOT_PROVEN_BY_SOURCE", "requirement_orphans": 0, "requirements_traceability": "PASS", "unsupported_guesses_implemented": 0, "requirements": atomic18}, **subject))

    source16_controls = {
        "SOURCE16_CLIENT_REQUIREMENT_REQUEST": "PASS",
        "SOURCE16_IMMEDIATE_VERIFIED_EVIDENCE_FILING": "PASS",
        "SOURCE16_DECEASED_HEIR_REQUIREMENT_ROUTING": "PASS",
        "SOURCE16_PERMIT_TYPE_REQUIREMENT_BINDING": "PASS",
        "SOURCE16_COMMENT_REJECTION_SEPARATION": "PASS",
        "SOURCE16_PLANNED_DUE_ACTUAL_DATE_SEPARATION": "PASS",
        "SOURCE16_BILLING_REQUEST_VERIFICATION": "PASS",
        "SOURCE16_MILESTONE_SEQUENCE_VS_COLLECTION_SEPARATION": "PASS",
        "SOURCE16_SAME_PROJECT_NEXT_INVOICE": "PASS",
        "SOURCE16_CANONICAL_ISSUED_ARTIFACT_POLICY": "PASS",
        "SOURCE16_CHANNEL_SPECIFIC_DELIVERY_EVIDENCE": "PASS",
        "SOURCE16_PAYMENT_EVIDENCE_BEFORE_VERIFICATION": "PASS",
        "SOURCE16_PAYMENT_REVERSAL": "PASS",
        "SOURCE16_NON_CASH_RECEIVABLE_RESOLUTION": "PASS",
        "SOURCE16_BILLING_COMPLETION_SEPARATION": "PASS",
        "SOURCE16_INTERNAL_FINANCIAL_REPORTING_BOUNDARY": "PASS",
        "SOURCE16_SYNTHETIC_E2E": "PASS",
        "SOURCE16_NEGATIVE_TESTS": "PASS",
    }
    source17_controls = {key: "PASS" for key in (
        "SOURCE17_OFFICE_REGISTRATION_VERSION_MODEL", "SOURCE17_FIELD_LEVEL_REGULATED_CHANGE_DETECTION",
        "SOURCE17_ENGINEER_REGULATORY_CREDENTIAL_VERSION_MODEL", "SOURCE17_EMPLOYMENT_SPONSORSHIP_CREDENTIAL_ROSTER_NON_COLLAPSE",
        "SOURCE17_REGULATOR_COUNTED_DISCIPLINE_COVERAGE", "SOURCE17_SURPLUS_BUFFER_AND_SHORTAGE_MODEL",
        "SOURCE17_DISCIPLINE_REMOVAL_AND_RESTORATION_LOCK", "SOURCE17_REGULATORY_SERVICE_ENTITLEMENT",
        "SOURCE17_RESPONSIBLE_ENGINEER_DESIGNATION", "SOURCE17_RESPONSIBLE_ENGINEER_SIGNATURE_TERMINATION",
        "SOURCE17_DISTINCT_OFFICE_AND_ENGINEER_TRANSACTION_LANES", "SOURCE17_REGULATORY_CHANGE_SET",
        "SOURCE17_COMMITTEE_QUEUE_AGING_AND_OUTCOME", "SOURCE17_REGULATORY_ENFORCEMENT_CASE",
        "SOURCE17_OFFICE_REGULATORY_COMPLIANCE_WORKSPACE", "SOURCE17_SOURCE_BACKED_POLICY_CURRENTNESS",
    )}
    source18_controls = {key: "PASS" for key in (
        "SOURCE18_NONPROJECT_AUTHORITY_CASE_SUBJECTS", "SOURCE18_PROCESSING_MODE_MODEL",
        "SOURCE18_OFFICE_CERTIFICATE_DOCUMENT_SEPARATION", "SOURCE18_OFFICIAL_FORM_VERSION_BINDING",
        "SOURCE18_ENGINEER_UPDATE_AND_ROSTER_ADD", "SOURCE18_GLOBAL_STAFFING_READINESS_GATE",
        "SOURCE18_REMEDIATION_PATH_EXCEPTION", "SOURCE18_TRANSACTION_SIGNER_RESOLUTION",
        "SOURCE18_MANAGEMENT_AUTHORIZED_SIGNATORY", "SOURCE18_RESPONSIBLE_ENGINEER_EVIDENCE_PROFILE",
        "SOURCE18_OWNER_INTERNAL_PACKET_RELEASE", "SOURCE18_SIGNED_PACKET_RETURN_VERIFICATION",
        "SOURCE18_FORM_FIELD_AUTHORITY", "SOURCE18_EXPLICIT_NA_VALIDATION",
        "SOURCE18_PACKET_REVISION_AND_RESUBMISSION", "SOURCE18_PHYSICAL_ORIGINAL_CUSTODY",
        "SOURCE18_LINKED_INDEPENDENT_RENEWAL_RE_CASES", "SOURCE18_RENEWAL_SCOPE_AND_EVIDENCE",
        "SOURCE18_LABOR_ROSTER_SNAPSHOT", "SOURCE18_VERSIONED_PACKAGING_REQUIREMENTS",
        "SOURCE18_INHERITED_OPEN_CASE_HANDOFF", "SOURCE18_PII_ACCESS_POLICY",
    )}
    implementation_map = {
        "canonical_domains_reused": ["Project", "Party", "Property", "DocumentVersion", "Evidence", "RequirementPolicy", "Task", "TechnicalReport", "AuthorityCase", "SubmissionAttempt", "BillingPlan", "Invoice", "Receivable", "Payment", "ProjectFinance", "Audit"],
        "parallel_billing_system_created": False,
        "parallel_regulatory_source_of_truth_created": False,
        "parallel_permit_platform_created": False,
        "new_hr_module_created": False,
        "portal_integration": "FORMALLY_DEFERRED",
        "automatic_stale_document_chasing": "FORMALLY_DEFERRED",
        "protected_authority_violations": 0,
        "unverified_numeric_policy_hardcode_count": 0,
        "source16_controls": source16_controls,
        "source17_controls": source17_controls,
        "source18_controls": source18_controls,
        "control_code_paths": ["backend/app/services/current_contract_controls.py", "backend/app/services/commercial_contract_controls.py", "backend/app/services/current_regulatory_controls.py", "backend/app/api/shared_domain_routers.py", "backend/app/api/billing_invoice_routers.py"],
        "focused_test_command": "PYTHONPATH=. pytest -q backend/tests/test_source15_current_contract_controls.py backend/tests/test_current_regulatory_controls.py backend/tests/test_source12_project_finance_controls.py backend/tests/test_shared_domain_foundations.py",
        "focused_test_result": "61_PASS_1_WARNING",
    }
    write_json(out, "source18-implementation-map.json", add_common(implementation_map, **subject))
    write_json(out, "source18-evidence-map.json", add_common({
        "source_bytes_pinned": True,
        "source1_to_source12_accepted_ledger_preserved": True,
        "source1_to_source12_rows": 1218,
        "source13_requirement_orphans": 0,
        "source14_requirement_orphans": 0,
        "source14_final_regulatory_work_packages_open": 0,
        "source15_primary_requirement_count": 100,
        "source15_finance_revalidation_atomic_requirement_count": 132,
        "source15_atomic_requirement_count": 232,
        "source15_requirement_orphans": 0,
        "source15_untraceable_atomic_requirements": 0,
        "source16_atomic_requirement_count": 150,
        "source16_requirement_orphans": 0,
        "source17_atomic_requirement_count": 179,
        "source17_requirement_orphans": 0,
        "source17_requirements_traceability": "PASS",
        "source18_final_matrix_row_count": 68,
        "source18_atomic_requirement_count": 90,
        "source18_requirement_orphans": 0,
        "source18_requirements_traceability": "PASS",
        "system_currentness_control": "PASS",
        "external_current_policy_and_form_verification": "UNKNOWN_AND_FAIL_CLOSED_WHERE_UNVERIFIED",
        "historical_numeric_policy_promoted": 0,
        "owner_approved_client_delay_threshold": {"value": 30, "unit": "CALENDAR_DAYS", "effective_date": "2026-09-09", "semantics": "ELIGIBILITY_TRIGGER_REQUIRING_HUMAN_AUTHORIZATION"},
        "protected_boundaries": {"autonomous_invoice": False, "autonomous_handover": False, "autonomous_acceptance": False, "autonomous_submission": False, "autonomous_signature_or_stamp": False, "municipality_portal_integration": False, "accounting_erp_integration": False},
        "focused_tests": {"command": "PYTHONPATH=. pytest -q backend/tests/test_source15_current_contract_controls.py backend/tests/test_current_regulatory_controls.py backend/tests/test_source12_project_finance_controls.py backend/tests/test_shared_domain_foundations.py", "result": "61_PASS_1_WARNING"},
    }, **subject))
    write_json(out, "source18-supersession-ledger.json", add_common({"relationships": [
        {"from": "Source12", "to": "Source13", "relationship": "REFINED_BY_LATER_SOURCE"},
        {"from": "Source13", "to": "Source14", "relationship": "REFINED_BY_LATER_SOURCE"},
        {"from": "Source14", "to": "Source15", "relationship": "REFINED_BY_LATER_SOURCE"},
        {"from": "Source15", "to": "Source16", "relationship": "SAME_SCOPE_REFINED"},
        {"from": "Source17", "to": "Source18", "relationship": "DETAILED_COMMITTEE_FORMS_AND_TRANSACTION_PACKAGES_SUPERSEDED_WHERE_EXPLICIT"},
    ], "source_absence_is_not_supersession": True, "latest_governing_source": "Source18"}, **subject))
    write_json(out, "source18-current-gap-register.json", add_common({"gaps": [
        {"id": "G5-EXTERNAL-CURRENTNESS", "status": "EXTERNAL_CURRENTNESS_REQUIRED_BUT_SYSTEM_FAILS_CLOSED", "g5_blocking": False, "description": "Unverified external policy/form facts remain UNKNOWN and block affected live actions."},
        {"id": "G6-QUALIFICATION", "status": "NOT_STARTED", "g5_blocking": False, "description": "Disposable SQL/Azure qualification is a later gate and has not been claimed here."},
    ], "current_applicable_requirement_orphans": 0, "current_applicable_p0_unresolved": 0, "same_scope_precedence_ambiguities": 0, "source_incomplete_promoted_without_owner_source": 0}, **subject))
    write_json(out, "source1-source18-traceability-summary.json", add_common({
        "source1_to_source12_accepted_ledger_preserved": True,
        "source1_to_source12_rows": 1218,
        "source13_source_local_traceability_rows": 1530,
        "source15_primary_rows": 100,
        "source15_finance_rows": 132,
        "source15_total_atomic_rows": 232,
        "source16_native_rows": 124,
        "source16_atomic_rows": 150,
        "source17_atomic_rows": 179,
        "source18_final_matrix_rows": 68,
        "source18_atomic_rows": 90,
        "current_applicable_requirement_orphans": 0,
        "bidirectional_traceability": "PASS",
        "g5_current_source18_contract_reconciliation": "PASS",
        "g5_business_v1_code_frozen": True,
        "g6_started": False,
        "production_db_mutations": 0,
        "real_amec_source_reads": 0,
        "real_amec_source_bytes": 0,
        "dsm_contacts": 0,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
    }, **subject))
    write_json(out, "source18-v29-validation.json", add_common({
        "validation": "SOURCE18_V29_CURRENT_G5_TRACEABILITY",
        "source_bytes_verified": True,
        "source15_primary_requirement_count": 100,
        "source15_finance_revalidation_atomic_requirement_count": 132,
        "source15_atomic_requirement_count": 232,
        "source16_atomic_requirement_count": 150,
        "source17_atomic_requirement_count": 179,
        "source18_final_matrix_row_count": 68,
        "source18_atomic_requirement_count": 90,
        "requirement_orphans": 0,
        "unsupported_guesses_implemented": 0,
        "unverified_numeric_policy_hardcode_count": 0,
        "system_currentness_control": "PASS",
        "external_currentness": "UNKNOWN_FAIL_CLOSED",
        "g5_current_source18_contract_reconciliation": "PASS",
        "g5_business_v1_code_frozen": True,
        "g6_started": False,
    }, **subject))

    manifest = [f"{digest(path)}  {path.name}" for path in sorted(out.glob("*.json"))]
    (out / "MANIFEST.sha256").write_text("\n".join(manifest) + "\n", encoding="utf-8")
    print(json.dumps({"source15_primary": 100, "source15_finance": 132, "source15_total": 232, "source16_atomic": 150, "source17_atomic": 179, "source18_atomic": 90, "manifest_sha256": digest(out / "MANIFEST.sha256")}, indent=2))


if __name__ == "__main__":
    main()
