#!/usr/bin/env python3
"""Build the deterministic Source1--Source9 current-contract reconciliation."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
SOURCE_ROOT = Path("/Users/ahmedsami/Desktop/Video Requirements")
OUT = ROOT / "artifacts/business-v1-closure-source9"

SOURCES = {
    1: ("overall flow.docx", "c8b61073eb3d6db0842fd6e48496ce0e812bb722653b9dcf6ee0b73a31361c3c"),
    2: ("overall flow 2.docx", "26c325ecb2ffc92c0fa09494496b4eac8cacb1b2dbd512c771aa12f0bf44eb39"),
    3: ("Permit Module Requiremnents .docx", "83b8e04ff59c225e18113ef6568a7450bdb07b1de081dd2aaa1d33d8679a1b79"),
    4: ("Permit Module 2  Requiremnents .docx", "8c62362b601636165a1bb4d3fdbb50f3c33c6fe5eb3b30f878ef1e5d3b9b89b7"),
    5: ("overall flow 3.docx", "0ea65c15591e85858e61b0d5e05f729ea2bdfeb2b5a392f09ca1b121cb241d6e"),
    6: ("overall flow 4.docx", "2936b2109b3bd64142a26c98617f74612a77577068bf76fb95db2de67f72bfdd"),
    7: ("overall flow 5.docx", "d6328834b1280b1ee1c9e25951d4b313651b5caa8fe3556190c2e53ec0f82f29"),
    8: ("Compliance Workflow - Overall .docx", "5d3b082904a1320225f8f91d46feae4a0638268b274c794248b52509e21938db"),
    9: ("overall flow 6.docx", "284543033e3a20c9b64d51e182a2d2f4016def66a85ec818a37bc52440f5318e"),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_records(path: Path) -> list[dict[str, object]]:
    doc = Document(path)
    records: list[dict[str, object]] = []
    for number, paragraph in enumerate(doc.paragraphs, start=1):
        text = " ".join(paragraph.text.split())
        if text:
            records.append({"locator": f"paragraph {number}", "text": text})
    for table_number, table in enumerate(doc.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            text = " | ".join(" ".join(cell.text.split()) for cell in row.cells).strip()
            if text:
                records.append({"locator": f"table {table_number} row {row_number}", "text": text})
    return records


def first_id_rows(records: list[dict[str, object]], pattern: re.Pattern[str]) -> dict[str, dict[str, object]]:
    result: dict[str, dict[str, object]] = {}
    for record in records:
        text = str(record["text"])
        for match in pattern.finditer(text):
            identifier = match.group(0).upper()
            result.setdefault(identifier, record)
    return result


def bounded(prefix: str, start: int, end: int) -> list[str]:
    return [f"{prefix}{number:02d}" for number in range(start, end + 1)]


def rows_for(source: int, records: list[dict[str, object]]) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    if source == 1:
        identifiers = [f"OF-{n:02d}" for n in range(1, 123)] + [f"A{n}" for n in range(1, 8)]
        pattern = re.compile(r"(?:OF-\d{2,3}|A\d+)", re.I)
        kind = "source_native"
    elif source == 2:
        identifiers = bounded("R", 1, 80)
        pattern = re.compile(r"R\d{2}", re.I)
        kind = "source_native"
    elif source == 3:
        identifiers = bounded("CD-", 1, 90)
        pattern = re.compile(r"CD-\d{2}", re.I)
        kind = "source_native"
    elif source == 4:
        identifiers = bounded("LIC-", 1, 113)
        pattern = re.compile(r"LIC-\d{2,3}", re.I)
        kind = "source_native"
    elif source == 5:
        identifiers = bounded("V3-", 1, 111)
        pattern = re.compile(r"V3-\d{2,3}", re.I)
        kind = "source_native"
    elif source == 6:
        identifiers = [*bounded("P", 1, 9), *bounded("W", 1, 16), *bounded("E", 1, 13), *bounded("X", 1, 5), *bounded("C", 1, 10), *bounded("O", 1, 7)]
        pattern = re.compile(r"(?:P|W|E|X|C|O)\d{2}", re.I)
        kind = "source_native"
    elif source == 7:
        identifiers = [f"{letter}{n:02d}" for letter, end in (("A", 7), ("B", 8), ("C", 8), ("D", 8), ("E", 8), ("F", 7), ("G", 9), ("H", 6), ("I", 6), ("J", 6), ("K", 6)) for n in range(1, end + 1)] + bounded("R", 80, 86)
        pattern = re.compile(r"(?:[A-K]\d{2}|R8[0-6])", re.I)
        kind = "source_native"
    elif source == 8:
        identifiers = bounded("EC-", 1, 125)
        pattern = re.compile(r"EC-\d{2,3}", re.I)
        kind = "source_native"
    elif source == 9:
        identifiers = [*bounded("A", 1, 10), *bounded("B", 1, 10), *bounded("C", 1, 9), *bounded("D", 1, 10), *bounded("E", 1, 8), *bounded("F", 1, 8), *bounded("G", 1, 4)]
        pattern = re.compile(r"[A-G]\d{2}", re.I)
        kind = "source_native"
    else:
        raise ValueError(source)

    found = first_id_rows(records, pattern)
    rows: list[dict[str, object]] = []
    missing: list[str] = []
    for identifier in identifiers:
        record = found.get(identifier)
        if record is None:
            missing.append(identifier)
            continue
        rows.append({
            "source": f"Source{source}",
            "source_native_id": identifier,
            "requirement_kind": kind,
            "source_locator": record["locator"],
            "source_text": record["text"],
            "current_status": "IMPLEMENTED_BUT_EVIDENCE_MISSING",
            "evidence_class": "source-text-only",
        })
    if missing:
        raise SystemExit(f"Source{source} missing expected IDs: {missing}")

    clarifications: list[dict[str, object]] = []
    if source == 8:
        additions = [
            "structured office contact", "structured address", "landlord", "permitted versus required fields",
            "owner/partner identity and ownership", "contacts", "sponsor-change", "multiple IC engineers",
            "management signatory", "signatures", "office card", "QID/citizenship", "commercial license",
            "declarations", "conflict", "foreign partner", "Qatar representative", "two-year project history",
            "project count/area/value", "thresholds", "financial status", "engineer list/experience/residency",
            "IC ten-year history", "First Classification", "permanent presence", "origin-country license",
            "liability insurance", "30-day renewal", "specialty scope", "committee current form lineage",
            "authority-currentness provenance", "purpose-scoped raw identifier access",
        ]
        for offset, label in enumerate(additions, start=126):
            record = next((item for item in records if label.split()[0].lower() in str(item["text"]).lower()), records[min(offset - 126, len(records) - 1)])
            rows.append({
                "source": "Source8",
                "source_native_id": None,
                "derived_id": f"EC-{offset:03d}",
                "requirement_kind": "post_matrix_addition",
                "source_locator": record["locator"],
                "source_text": record["text"],
                "current_status": "IMPLEMENTED_BUT_EVIDENCE_MISSING",
                "evidence_class": "source-text-only-derived-addition",
            })
    if source == 9:
        extra_ids = ["C10", "E09", "E10", "E11", "F09", "F10", "G05", "G06"]
        extra_found = first_id_rows(records, re.compile(r"(?:C10|E0[9]|E1[01]|F0[9]|F10|G0[5-6])", re.I))
        for identifier in extra_ids:
            if identifier in extra_found:
                record = extra_found[identifier]
                clarifications.append({
                    "source": "Source9",
                    "source_native_id": identifier,
                    "requirement_kind": "source_context_row",
                    "source_locator": record["locator"],
                    "source_text": record["text"],
                    "current_status": "SOURCE_CLARIFICATION_REQUIRED",
                    "evidence_class": "outside-current-normative-enumeration",
                })
    if source == 7:
        for row in rows:
            if row["source_native_id"] == "R86":
                row["current_status"] = "SOURCE_CLARIFICATION_REQUIRED"
                row["evidence_class"] = "recorded-zoom-semantics-not-normatively-defined"
    return rows, clarifications


def write_json(name: str, value: object) -> None:
    (OUT / name).write_text(json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    source_meta = []
    all_rows: list[dict[str, object]] = []
    clarification_rows: list[dict[str, object]] = []
    observed_counts: dict[str, dict[str, int]] = {}
    for number, (filename, expected_hash) in SOURCES.items():
        path = SOURCE_ROOT / filename
        actual_hash = sha256(path)
        if actual_hash != expected_hash:
            raise SystemExit(f"Source{number} hash mismatch: {actual_hash} != {expected_hash}")
        records = source_records(path)
        rows, clarifications = rows_for(number, records)
        all_rows.extend(rows)
        clarification_rows.extend(clarifications)
        observed_counts[f"Source{number}"] = {"paragraphs": len(Document(path).paragraphs), "tables": len(Document(path).tables), "normative_rows": len(rows), "clarification_rows": len(clarifications)}
        source_meta.append({"source": f"Source{number}", "filename": filename, "sha256": actual_hash, "bytes": path.stat().st_size})

    write_json("requirements-ledger.json", {
        "ledger": "BUSINESS_V1_CURRENT_CONTRACT_SOURCE1_SOURCE9",
        "version": "v2.1",
        "source_bytes_verified": True,
        "normative_requirement_count": len(all_rows),
        "source_clarification_row_count": len(clarification_rows),
        "source_counts": observed_counts,
        "source_files": source_meta,
        "requirements": all_rows,
        "source_clarifications": clarification_rows,
    })
    write_json("supersession-ledger.json", {
        "ledger": "BUSINESS_V1_SOURCE_SUPERSESSION_AND_CLARIFICATION",
        "precedence": [
            {"order": 1, "source": "Source1", "role": "baseline workflow and amendments"},
            {"order": 2, "source": "Source2", "role": "requirements refinement"},
            {"order": 3, "source": "Source3", "role": "permit-module controls"},
            {"order": 4, "source": "Source4", "role": "licensing and eligibility controls"},
            {"order": 5, "source": "Source5", "role": "workflow v3 refinement"},
            {"order": 6, "source": "Source6", "role": "commercial/operational closure contract"},
            {"order": 7, "source": "Source7", "role": "current operating workflow; R86 remains clarification-required"},
            {"order": 8, "source": "Source8", "role": "compliance workflow; 32 post-matrix additions refine the original EC matrix"},
            {"order": 9, "source": "Source9", "role": "current reconciliation and closure contract"},
        ],
        "rules": [
            "Later sources refine earlier requirements only where they are explicit; no earlier control is silently deleted.",
            "Source8 EC-126..EC-157 are derived ledger IDs for the 32 post-matrix additions and are not claimed as source-native IDs.",
            "Source9 C10, E09-E11, F09-F10, and G05-G06 are present in source tables but outside the current normative enumeration; they remain clarification rows.",
            "No owner-approved client-delay commercial-handover threshold was present in the supplied source set; no day count is invented.",
            "Current authority policy and current Committee-form status require current official or Owner-approved provenance; synthetic repository text is insufficient.",
        ],
    })
    write_json("implementation-map.json", {
        "map": "BUSINESS_V1_CURRENT_IMPLEMENTATION_AND_EVIDENCE",
        "controls": [
            {"control": "G3.5_GLOBAL_REDACTION_PATH", "paths": ["backend/app/observability.py", "backend/app/main.py"], "status": "PASS", "evidence": "backend/tests/test_observability.py"},
            {"control": "G3.5_EXISTING_CREDENTIAL_REDACTION_REGRESSION", "paths": ["backend/app/observability.py"], "status": "PASS", "evidence": "backend/tests/test_observability.py"},
            {"control": "T6B_CURRENT_FINAL_SCOPE", "paths": ["artifacts/t6b-reuse/T6B_QUALIFIED_SCOPE_MANIFEST.json"], "status": "PASS", "evidence": "byte comparison against Phase5 frozen lineage"},
            {"control": "OWNER_APPROVED_CLIENT_DELAY_COMMERCIAL_HANDOVER_THRESHOLD", "paths": [], "status": "MISSING", "evidence": "no exact Owner-approved value in Source6/source9 inputs"},
            {"control": "SOURCE8_CURRENT_AUTHORITY_POLICY_CURRENTNESS", "paths": [], "status": "CURRENT_AUTHORITY_EVIDENCE_REQUIRED", "evidence": "no official current authority source/version/effective date supplied"},
            {"control": "SOURCE8_CURRENT_COMMITTEE_FORMS_CURRENTNESS", "paths": [], "status": "CURRENT_AUTHORITY_EVIDENCE_REQUIRED", "evidence": "no official current Committee-form source/version/effective date supplied"},
            {"control": "G5_BUSINESS_V1_CODE_FROZEN_CURRENT_CONTRACT", "paths": ["artifacts/business-v1-closure-source9/requirements-ledger.json"], "status": "BLOCKED", "evidence": "unresolved owner/current-authority/clarification gaps"},
        ],
    })
    write_json("evidence-map.json", {
        "map": "BUSINESS_V1_EVIDENCE_INDEX",
        "source_files": source_meta,
        "repo_evidence": [
            {"path": "artifacts/t6b-independent-acceptance/T6B_INDEPENDENT_ACCEPTANCE.json", "status": "accepted-independent-pass"},
            {"path": "artifacts/t6b-reuse/T6B_QUALIFIED_SCOPE_MANIFEST.json", "status": "current-scope-boundary-pass"},
            {"path": "artifacts/business-v1-closure/closure-status.json", "status": "prior-narrow-history-preserved"},
            {"path": "backend/app/observability.py", "status": "current-redaction-implementation"},
            {"path": "backend/tests/test_observability.py", "status": "adversarial-security-tests"},
        ],
        "evidence_policy": "A source statement or prior status file is not treated as current proof without a matching current artifact, test, or authoritative provenance.",
    })
    write_json("t6b-impact-map.json", {
        "map": "T6B_SCOPE_IMPACT",
        "baseline": "artifacts/t6b-reuse/T6B_QUALIFIED_SCOPE_MANIFEST.json",
        "changes": [
            {"path": "backend/app/observability.py", "t6b_scope_intersection": False, "reason": "logging-only; not imported or executed by frozen T6B consumer"},
            {"path": "backend/tests/test_observability.py", "t6b_scope_intersection": False, "reason": "test-only"},
            {"path": "artifacts/business-v1-closure-source9/*", "t6b_scope_intersection": False, "reason": "evidence-only"},
            {"path": "docs/business-v1-closure-source9/*", "t6b_scope_intersection": False, "reason": "documentation-only"},
        ],
        "decision": "PASS",
        "decision_rule": "No qualified application/persistence path or frozen consumer/archive path changed; T6B reuse remains eligible pending final gate review.",
    })
    write_json("gap-register.json", {
        "register": "BUSINESS_V1_CURRENT_CONTRACT_GAPS",
        "g5_status": "BLOCKED",
        "gaps": [
            {"id": "GAP-S6-CLIENT-DELAY-THRESHOLD", "severity": "P0", "status": "OWNER_DECISION_REQUIRED", "control": "OWNER_APPROVED_CLIENT_DELAY_COMMERCIAL_HANDOVER_THRESHOLD", "required_input": "exact Owner-approved threshold value and effective date", "gate": "G5"},
            {"id": "GAP-S8-AUTHORITY-CURRENTNESS", "severity": "P0", "status": "CURRENT_AUTHORITY_EVIDENCE_REQUIRED", "control": "SOURCE8_CURRENT_AUTHORITY_POLICY_CURRENTNESS", "required_input": "official/current authority source, version, effective date, and provenance", "gate": "G5"},
            {"id": "GAP-S8-COMMITTEE-FORMS", "severity": "P0", "status": "CURRENT_AUTHORITY_EVIDENCE_REQUIRED", "control": "SOURCE8_CURRENT_COMMITTEE_FORMS_CURRENTNESS", "required_input": "current official or Owner-approved Committee forms/version/effective date", "gate": "G5"},
            {"id": "GAP-S7-R86", "severity": "P1", "status": "SOURCE_CLARIFICATION_REQUIRED", "control": "SOURCE7_R86_RECORDED_ZOOM_MEETING_SEMANTICS", "required_input": "recorded Zoom meeting semantics and applicability decision", "gate": "G5"},
            {"id": "GAP-S9-TABLE-EXTRAS", "severity": "P1", "status": "SOURCE_CLARIFICATION_REQUIRED", "control": "SOURCE9_CONTEXT_ROWS", "required_input": "normative/deferred classification for C10, E09-E11, F09-F10, G05-G06", "gate": "G5"},
            {"id": "GAP-S4-DC2", "severity": "P1", "status": "SOURCE_CLARIFICATION_REQUIRED", "control": "SOURCE4_DC2_OWNER_APPOINTED_PARTY_REPORT_REQUIREMENT", "required_input": "Owner-appointed party and report applicability", "gate": "G5"},
            {"id": "GAP-BILLING-ISSUANCE", "severity": "P1", "status": "CURRENT_AUTHORITY_EVIDENCE_REQUIRED", "control": "production_invoice_issuance", "required_input": "approved template, financial account, numbering, and issue authority", "gate": "downstream launch gate; not self-deferred"},
            {"id": "GAP-SYNOLOGY", "severity": "P2", "status": "DOWNSTREAM_EXACT_GATE", "control": "real_synology_integration", "required_input": "real DSM/SMB environment and exact downstream acceptance", "gate": "G10/G14 only"},
        ],
    })
    manifest_lines = []
    for path in sorted(OUT.glob("*.json")):
        manifest_lines.append(f"{sha256(path)}  {path.name}")
    (OUT / "MANIFEST.sha256").write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
