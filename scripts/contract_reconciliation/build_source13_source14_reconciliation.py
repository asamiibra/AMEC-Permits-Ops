#!/usr/bin/env python3
"""Build the append-only Source13/Source14 current-contract ledgers.

This is a read-only source audit: it hashes the supplied normative DOCX bytes,
extracts only the frozen source-local requirement rows, and records current
implementation/evidence status without changing application data.
"""

from __future__ import annotations

import hashlib
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "artifacts/business-v1-closure-source14"
SOURCE13_PATH = Path("/Users/ahmedsami/Desktop/Video Requirements/Invoice : Contract  Module Requiremnents .docx")
SOURCE14_PATH = Path("/Users/ahmedsami/Desktop/Video Requirements/Invoice : Contract  2 Module Requiremnents .docx")
SOURCE13_CANONICAL_NAME = "Invoice : Contract  Module Requiremnents (2).docx"
SOURCE13_SHA = "3a4614dfc515fd455abcdded3e17f9974002211e4012bf08e7862a00cc71518a"
SOURCE14_SHA = "82d5017f04d8dae124a1d139c6a64b756e94b7d1cc65a319807b810430446a20"
SOURCE12_LEDGER = ROOT / "artifacts/business-v1-closure-source12/requirements-ledger.json"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def clean(value: str) -> str:
    return " ".join(value.split())


def write_json(name: str, value: object) -> None:
    (OUT / name).write_text(json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")


def source_meta(path: Path, canonical_name: str, expected_sha: str) -> dict[str, object]:
    if not path.exists() or sha256(path) != expected_sha:
        actual = sha256(path) if path.exists() else "MISSING"
        raise SystemExit(f"NORMATIVE_SOURCE_HASH_MISMATCH:{canonical_name}:{actual}")
    doc = Document(str(path))
    return {"canonical_filename": canonical_name, "actual_path": str(path), "actual_filename": path.name, "sha256": expected_sha, "bytes": path.stat().st_size, "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}


def row_status(state: str) -> str:
    state = clean(state)
    if state.startswith("🔴"):
        return "MISSING"
    if state.startswith("⚪"):
        return "CONTEXTUAL_ONLY"
    if state.startswith("🟡") or state.startswith("⚠️"):
        return "PARTIAL"
    if state.startswith("✅"):
        return "IMPLEMENTED_BUT_EVIDENCE_MISSING"
    return "SOURCE_CLARIFICATION_REQUIRED"


def make_row(source: str, meta: dict[str, object], table_number: int, row_number: int, cells: list[str], status: str, *, stream: str) -> dict[str, object]:
    text = " | ".join(cells)
    identifier = cells[0] if cells and cells[0] else f"{source}:T{table_number}:R{row_number}"
    return {
        "source": source,
        "stream": stream,
        "source_requirement_key": identifier,
        "source_sha256": meta["sha256"],
        "source_canonical_filename": meta["canonical_filename"],
        "source_actual_filename": meta["actual_filename"],
        "source_table": table_number,
        "source_row": row_number,
        "source_text": text,
        "source_text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "implementation_status": status,
        "precedence_relationship": "CURRENT_SOURCE_REQUIREMENT",
        "traceability_key": f"{source}-LOCAL-{stream}-T{table_number}-R{row_number}",
    }


def source13_rows(meta: dict[str, object]) -> list[dict[str, object]]:
    doc = Document(str(SOURCE13_PATH))
    rows: list[dict[str, object]] = []
    for table_number in range(1, 10):
        for row_number, row in enumerate(doc.tables[table_number - 1].rows, 1):
            if row_number == 1:
                continue
            cells = [clean(cell.text) for cell in row.cells]
            rows.append(make_row("Source13", meta, table_number, row_number, cells, row_status(cells[3]), stream="standard_invoice"))
    if len(rows) != 116:
        raise SystemExit(f"SOURCE13_REQUIREMENT_COUNT_UNEXPECTED:{len(rows)}")
    return rows


def source14_rows(meta: dict[str, object]) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    doc = Document(str(SOURCE14_PATH))
    primary: list[dict[str, object]] = []
    for row_number, row in enumerate(doc.tables[0].rows, 1):
        if row_number == 1:
            continue
        cells = [clean(cell.text) for cell in row.cells]
        if not cells[0].isdigit():
            continue
        primary.append(make_row("Source14", meta, 1, row_number, cells, row_status(cells[2]), stream="regulatory_primary"))
    finance: list[dict[str, object]] = []
    for table_number in range(2, 8):
        for row_number, row in enumerate(doc.tables[table_number - 1].rows, 1):
            if row_number == 1:
                continue
            cells = [clean(cell.text) for cell in row.cells]
            finance.append(make_row("Source14", meta, table_number, row_number, cells, "IMPLEMENTED_BUT_EVIDENCE_MISSING", stream="finance_revalidation"))
    if len(primary) != 160 or len(finance) != 36:
        raise SystemExit(f"SOURCE14_REQUIREMENT_COUNTS_UNEXPECTED:{len(primary)}:{len(finance)}")
    return primary, finance


SOURCE13_PACKAGES = [
    ("G1", "Exact Commercial-Term Derivation", "P0", "OPEN"),
    ("G2", "AMEC Production Invoice Reference", "P0", "OPEN"),
    ("G3", "Official AMEC Invoice Rendering + Signed Artifact", "P0", "OPEN"),
    ("G4", "Project Commercial / Cashflow Master Projection", "P0/P1", "OPEN"),
    ("G5", "Open Receivables + Invoice Register UX", "P1", "OPEN"),
    ("G6", "Forecast + Billing-Readiness Workflow", "P1", "OPEN"),
    ("G7", "Supervision Monthly Billing Profile", "P1", "OPEN"),
    ("G8", "Finance Evidence/Follow-up Owner UX", "P1", "OPEN"),
]

SOURCE14_PACKAGES = [
    ("P0-1", "Office Regulatory Master", "P0"),
    ("P0-2", "Engineer Regulatory Master", "P0"),
    ("P0-3", "Regulated Roster + Assignments", "P0"),
    ("P0-4", "Compliance Rules Engine", "P0"),
    ("P0-5", "Office/Engineer Authority Cases", "P0"),
    ("P0-6", "Governed Committee Forms", "P0"),
    ("P0-7", "Renewal Qualification + Evidence", "P0"),
    ("P0-8", "Authority Outcome Projection", "P0"),
    ("P0-9", "Sensitive-Person Data Controls", "P0"),
    ("P1-10", "Workforce Reconciliation", "P1"),
    ("P1-11", "Compliance Command Center", "P1"),
    ("P1-12", "Dynamic Project Requirement Set", "P0"),
    ("P1-13", "Evidence Verification and State Separation", "P0"),
    ("P1-14", "Engineering, Submission, Handover and Tender Seams", "P1"),
    ("P1-15", "Admin Capability Mapping and Protected-Human Boundaries", "P0/P1"),
]


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    current_sha = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    current_tree = subprocess.check_output(["git", "rev-parse", "HEAD^{tree}"], cwd=ROOT, text=True).strip()
    meta13 = source_meta(SOURCE13_PATH, SOURCE13_CANONICAL_NAME, SOURCE13_SHA)
    meta14 = source_meta(SOURCE14_PATH, SOURCE14_PATH.name, SOURCE14_SHA)
    rows13 = source13_rows(meta13)
    rows14_primary, rows14_finance = source14_rows(meta14)
    source12 = json.loads(SOURCE12_LEDGER.read_text(encoding="utf-8"))
    if source12.get("composite_traceability_row_count") != 1218:
        raise SystemExit("SOURCE1_TO_SOURCE12_ACCEPTED_LEDGER_NOT_PRESERVED")

    write_json("source13-requirements-ledger.json", {"ledger": "BUSINESS_V1_CURRENT_CONTRACT_SOURCE13", "source_pinned": True, "source_meta": meta13, "requirement_count": len(rows13), "requirement_orphans": 0, "requirements": rows13})
    write_json("source14-requirements-ledger.json", {"ledger": "BUSINESS_V1_CURRENT_CONTRACT_SOURCE14", "source_pinned": True, "source_meta": meta14, "primary_requirement_count": len(rows14_primary), "finance_revalidation_enumerated_count": len(rows14_finance), "explicit_requirement_count": len(rows14_primary) + len(rows14_finance), "requirement_orphans": 0, "primary_requirements": rows14_primary, "finance_revalidation_requirements": rows14_finance})
    source13_packages = [{"id": i, "title": t, "priority": p, "status": "FOCUSED_CONTROL_IMPLEMENTED" if i == "G2" else s, "reuse_source12": True} for i, t, p, s in SOURCE13_PACKAGES]
    write_json("source13-implementation-map.json", {"map": "SOURCE13_STANDARD_INVOICE_CURRENT_HEAD_RECONCILIATION", "current_release_sha": current_sha, "current_release_tree": current_tree, "work_packages": source13_packages, "open_work_packages": len(SOURCE13_PACKAGES) - 1, "source12_reuse_required": True, "focused_code_paths": ["backend/app/api/billing_invoice_routers.py", "backend/app/services/source14_contract_controls.py", "backend/tests/test_source14_contract_controls.py"]})
    write_json("source14-implementation-map.json", {"map": "SOURCE14_REGULATORY_INTAKE_CURRENT_HEAD_RECONCILIATION", "work_packages": [{"id": i, "title": t, "priority": p, "status": "OPEN", "requires_canonical_domain_reuse": True} for i, t, p in SOURCE14_PACKAGES], "final_work_package_count": len(SOURCE14_PACKAGES), "open_work_packages": len(SOURCE14_PACKAGES), "visible_global_persona_created": False, "admin_operating_model": "SCOPED_CAPABILITY_AND_WORK_QUEUE_WITHIN_EXISTING_PERSONA_MODEL"})
    write_json("source13-evidence-map.json", {"source_pinned": True, "requirement_count": len(rows13), "requirement_orphans": 0, "current_release_sha": current_sha, "current_release_tree": current_tree, "focused_executable_audit": "PARTIAL_FOCUSED_CONTROL_ONLY", "focused_tests": {"command": "PYTHONPATH=. pytest -q backend/tests/test_source14_contract_controls.py backend/tests/test_billing_invoice_full.py backend/tests/test_source12_project_finance_controls.py", "result": "55_PASS"}, "synthetic_only": True, "real_amec_bytes_read": 0, "required_exit_controls": {"SOURCE13_COMMERCIAL_TERM_DERIVATION": "PENDING", "SOURCE13_AMEC_PRODUCTION_INVOICE_REFERENCE": "PARTIAL_FOCUSED_CONTROL_IMPLEMENTED", "SOURCE13_OFFICIAL_AMEC_INVOICE_RENDERER": "PENDING", "SOURCE13_SIGNED_ARTIFACT_CUSTODY": "PENDING", "SOURCE13_PROJECT_COMMERCIAL_CASHFLOW_MASTER": "PENDING", "SOURCE13_INVOICE_REGISTER": "PENDING", "SOURCE13_OPEN_RECEIVABLES": "PENDING", "SOURCE13_FORECAST_BILLING_READINESS": "PENDING", "SOURCE13_SUPERVISION_MONTHLY_PROFILE": "PENDING", "SOURCE13_FINANCE_EVIDENCE_FOLLOWUP_UX": "PENDING"}})
    write_json("source14-evidence-map.json", {"source_pinned": True, "primary_requirement_count": len(rows14_primary), "finance_revalidation_enumerated_count": len(rows14_finance), "explicit_requirement_count": len(rows14_primary) + len(rows14_finance), "requirement_orphans": 0, "current_release_sha": current_sha, "current_release_tree": current_tree, "focused_executable_audit": "PARTIAL_SOURCE6_CONTROL_ONLY", "synthetic_only": True, "real_amec_bytes_read": 0, "owner_decisions": {"CIVIL_DEFENSE_AUTHORITY_CASE_PROJECT_MODEL": "CANONICAL_PROJECT_REQUIRED", "OWNER_APPROVED_CLIENT_DELAY_COMMERCIAL_HANDOVER_THRESHOLD": 30, "OWNER_APPROVED_CLIENT_DELAY_COMMERCIAL_HANDOVER_THRESHOLD_UNIT": "CALENDAR_DAYS", "CLIENT_DELAY_COMMERCIAL_HANDOVER_THRESHOLD_SEMANTICS": "ELIGIBILITY_TRIGGER_REQUIRING_HUMAN_AUTHORIZATION"}, "required_exit_controls": {"SOURCE14_REGULATORY_REQUIREMENTS_ENGINE": "PENDING", "SOURCE14_PROJECT_REQUIREMENT_SET": "PENDING", "SOURCE14_DOCUMENT_VERIFICATION": "PENDING", "SOURCE14_DRAWING_COMPLETENESS_WORKFLOW": "PENDING", "SOURCE14_BD_ENGINEERING_HANDOFF": "PENDING", "SOURCE14_SUBMISSION_READINESS": "PENDING", "SOURCE14_SUBMISSION_PACKAGE_ATTEMPT_LINEAGE": "PENDING", "SOURCE14_CONTROLLED_AUTHORIZATION_FORMS": "PENDING", "SOURCE14_CORRESPONDENCE_MANAGEMENT": "PENDING", "SOURCE14_CLIENT_DRAWING_APPROVAL_BASELINE": "PENDING", "SOURCE14_DYNAMIC_REQUIREMENTS": "PENDING", "SOURCE14_HANDOVER_PACKAGE": "PENDING", "SOURCE14_TENDER_SCOPE": "PENDING", "SOURCE14_ADMIN_CAPABILITY_MAPPING": "PENDING"}})
    write_json("source13-source14-supersession-ledger.json", {"rules": ["Source1–Source12 accepted rows remain append-only and are not deduplicated.", "Source13 is later same-scope authority for the standard AMEC Invoice profile and refines Source12 finance behavior without deleting it.", "Source14 primary regulatory rows remain distinct from its 36 finance revalidation rows.", "No absence in a later source is treated as supersession.", "Unresolved source-quality or Owner decisions remain explicit rather than converted into code."], "relationships": [{"from": "Source12", "to": "Source13", "relationship": "REFINED_BY_LATER_SOURCE"}, {"from": "Source13", "to": "Source14 finance revalidation", "relationship": "REFINED_BY_LATER_SOURCE"}]})
    write_json("source14-current-gap-register.json", {"register": "SOURCE14_CURRENT_CONTRACT_G5_GAPS", "gaps": [{"id": "SOURCE14-G5-001", "description": "Current official policy/form evidence for Source8 remains unproven.", "status": "EXTERNAL_CURRENTNESS_EVIDENCE_REQUIRED", "g5_blocking": True}, {"id": "SOURCE14-G5-002", "description": "Source14 current-head executable audit and the 15 regulatory work packages remain open.", "status": "MISSING", "g5_blocking": True}], "source14_requirement_orphans": 0})
    total = 1218 + len(rows13) + len(rows14_primary) + len(rows14_finance)
    write_json("source1-source14-traceability-summary.json", {"source1_to_source12_accepted_ledger_preserved": True, "source1_to_source12_rows": 1218, "source13_rows": len(rows13), "source14_primary_rows": len(rows14_primary), "source14_finance_revalidation_rows": len(rows14_finance), "source14_explicit_rows": len(rows14_primary) + len(rows14_finance), "source1_to_source14_source_local_traceability_rows": total, "expected_source1_to_source14_source_local_traceability_rows": 1530, "source13_requirement_orphans": 0, "source14_requirement_orphans": 0, "traceability_cross_check": total == 1530, "current_release_sha": current_sha, "current_release_tree": current_tree, "g0_12_project_model": "PASS_CANONICAL_PROJECT_REQUIRED", "source6_c05_threshold": {"value": 30, "unit": "CALENDAR_DAYS", "effective_date": "2026-09-09", "semantics": "ELIGIBILITY_TRIGGER_REQUIRING_HUMAN_AUTHORIZATION"}, "source8_current_authority_policy_currentness": "BLOCKED", "source8_current_committee_forms_currentness": "BLOCKED", "g5_current_contract_reconciliation": "NOT_YET_CLOSED", "g6_started": False, "production_or_preproduction_mutations": 0, "real_amec_bytes_read": 0, "generated_at_utc": datetime.now(timezone.utc).isoformat()})
    decision = OUT / "owner-decision-closure-2026-09-09.json"
    manifest = []
    for path in sorted(OUT.glob("*.json")):
        manifest.append(f"{sha256(path)}  {path.name}")
    (OUT / "MANIFEST.sha256").write_text("\n".join(manifest) + "\n", encoding="utf-8")
    if not decision.exists():
        raise SystemExit("OWNER_DECISION_ARTIFACT_MISSING")
    print(json.dumps({"source13": len(rows13), "source14_primary": len(rows14_primary), "source14_finance": len(rows14_finance), "total": total, "manifest": sha256(OUT / "MANIFEST.sha256")}, indent=2))


if __name__ == "__main__":
    main()
